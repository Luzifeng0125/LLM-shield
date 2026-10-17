"""Regression for WPS validation metadata and independently selectable sheets."""
from __future__ import annotations

import io
from zipfile import ZipFile

from fastapi.testclient import TestClient
from openpyxl import Workbook
from openpyxl.worksheet.datavalidation import DataValidation

from llm_shield.m1_access import file_loader
from llm_shield.server import app


def _workbook_with_wps_validation_id() -> bytes:
    workbook = Workbook()
    first = workbook.active
    first.title = "名单"
    first.merge_cells("A1:C1")
    first["A1"] = "合成标题"
    first.append(["姓名", "电话", "电话"])
    first.append(["人员甲", "13800138000", ""])
    validation = DataValidation(type="list", formula1='"是,否"')
    first.add_data_validation(validation)
    validation.add("A3")
    second = workbook.create_sheet("汇总")
    second.append(["学院", "人数"])
    second.append(["甲学院", 1])
    workbook.create_sheet("空白")
    original = io.BytesIO()
    workbook.save(original)
    modified = io.BytesIO()
    with ZipFile(io.BytesIO(original.getvalue())) as source, ZipFile(modified, "w") as target:
        for item in source.infolist():
            payload = source.read(item)
            if item.filename == "xl/worksheets/sheet1.xml":
                assert b"<dataValidation " in payload
                payload = payload.replace(b"<dataValidation ",
                                          b'<dataValidation id="wps-id" ', 1)
            target.writestr(item, payload)
    return modified.getvalue()


def test_excel_sheets_survive_wps_metadata_and_keep_headers(tmp_path, monkeypatch):
    monkeypatch.setattr(file_loader, "UPLOAD_DIR", tmp_path)
    materials = file_loader.parse_uploads(_workbook_with_wps_validation_id(), "合成.xlsx")
    assert len(materials) == 2
    assert materials[0].columns[:3] == ["姓名", "电话", "电话_2"]
    assert len(materials[0].rows) == 1
    assert materials[0].rows[0]["姓名"] == "人员甲"
    assert materials[1].columns[:2] == ["学院", "人数"]
    assert len(materials[1].rows) == 1
    assert len(list(tmp_path.glob("*.json"))) == 2
    assert not list(tmp_path.glob("*.xlsx"))


def test_upload_api_lists_all_nonempty_sheets(tmp_path, monkeypatch):
    monkeypatch.setattr(file_loader, "UPLOAD_DIR", tmp_path)
    with TestClient(app) as client:
        response = client.post("/api/upload", files={
            "file": ("合成.xlsx", _workbook_with_wps_validation_id(),
                     "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert response.status_code == 200
    body = response.json()
    assert body["worksheets"] == 2
    assert len(body["documents"]) == 2
    assert body["doc_id"] == body["documents"][0]["doc_id"]
