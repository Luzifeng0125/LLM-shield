"""Create an editable Chinese operating manual from the maintained Markdown."""
from pathlib import Path
import re
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

ROOT=Path(__file__).resolve().parents[1]
doc=Document();sec=doc.sections[0]
sec.page_width=Cm(21.59);sec.page_height=Cm(27.94)
sec.top_margin=Cm(1.8);sec.bottom_margin=Cm(1.8)
sec.left_margin=Cm(2.0);sec.right_margin=Cm(2.0)
sec.header_distance=Cm(.7);sec.footer_distance=Cm(.8)
for name,size,bold in [('Normal',11,False),('Title',25,True),('Heading 1',17,True),('Heading 2',13,True)]:
    style=doc.styles[name];style.font.name='Calibri';style.font.size=Pt(size);style.font.bold=bold
    style.font.color.rgb=RGBColor(0,0,0)
    style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'微软雅黑' if bold else '宋体')
    style.paragraph_format.line_spacing=1.22;style.paragraph_format.space_after=Pt(7)
    if name.startswith('Heading'):
        style.paragraph_format.space_before=Pt(13);style.paragraph_format.keep_with_next=True
for style in doc.styles:
    for b in list(style.element.iter(qn('w:pBdr'))):b.getparent().remove(b)
header=sec.header.paragraphs[0];header.text='隐盾项目  数据复核与训练操作手册'
for run in header.runs:run.font.size=Pt(9);run.font.color.rgb=RGBColor(0,0,0)
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run('2026 年 10 月 14 日   |   ')
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
for run in footer.runs:run.font.size=Pt(9)
doc.core_properties.title='隐盾数据复核与模型训练执行方案'
doc.core_properties.author='隐盾项目组'
doc.core_properties.subject='基于隐盾数据集 v1 的双人标注 裁决 人员培训与模型训练'


def table(lines):
    rows=[[c.strip() for c in line.strip().strip('|').split('|')] for line in lines]
    rows=[row for row in rows if not all(re.fullmatch(r':?-+:?',v) for v in row)]
    cols=len(rows[0]);widths={2:[5.2,12.39],3:[4.2,8.7,4.69],4:[2.4,4.8,5.0,5.39],5:[3.8,2,2,2,7.79],6:[3.5,2.1,2.1,2.7,2.6,4.59]}[cols]
    if rows[0][0]=='工作':widths=[5.0,3.9,8.69]
    if rows[0][0]=='试标 ID':widths=[3.0,4.8,9.79]
    if rows[0][0]=='动作':widths=[4.5,6.5,6.59]
    tab=doc.add_table(rows=0,cols=cols);tab.alignment=WD_TABLE_ALIGNMENT.CENTER;tab.autofit=False
    for col,w in zip(tab.columns,widths):col.width=Cm(w)
    for i,row in enumerate(rows):
        rr=tab.add_row();trpr=rr._tr.get_or_add_trPr();trpr.append(OxmlElement('w:cantSplit'))
        if i==0:trpr.append(OxmlElement('w:tblHeader'))
        for j,(cell,text,w) in enumerate(zip(rr.cells,row,widths)):
            cell.width=Cm(w);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell.text=text;cp=cell._tc.get_or_add_tcPr()
            borders=OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right','insideH','insideV']:
                e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
            cp.append(borders)
            mar=OxmlElement('w:tcMar')
            for edge,value in [('top','90'),('bottom','90'),('left','110'),('right','110')]:
                e=OxmlElement('w:'+edge);e.set(qn('w:w'),value);e.set(qn('w:type'),'dxa');mar.append(e)
            cp.append(mar)
            if i==0:
                shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E8EFF6');cp.append(shade)
            for p in cell.paragraphs:
                p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1.15
                if cols>=5 and j in [1,2,3]:p.alignment=WD_ALIGN_PARAGRAPH.CENTER
                for run in p.runs:run.font.size=Pt(10.5);run.bold=(i==0)
    doc.add_paragraph().paragraph_format.space_after=Pt(0)


lines=(ROOT/'隐盾数据复核与模型训练执行方案.md').read_text(encoding='utf-8').splitlines()
i=0
while i<len(lines):
    line=lines[i].strip()
    if not line:i+=1;continue
    if line.startswith('|'):
        group=[]
        while i<len(lines) and lines[i].strip().startswith('|'):group.append(lines[i]);i+=1
        table(group);continue
    if line.startswith('# '):doc.add_paragraph(line[2:],'Title')
    elif line.startswith('## '):doc.add_heading(line[3:],1)
    elif line.startswith('### '):doc.add_heading(line[4:],2)
    else:doc.add_paragraph(line)
    i+=1
out=ROOT/'隐盾数据复核与模型训练执行方案.docx';doc.save(out)
print(str(out))
