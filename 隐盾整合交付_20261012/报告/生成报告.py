"""从随包验证记录生成可编辑项目报告。运行：runtime/python/python.exe 报告/生成报告.py"""
import json
from pathlib import Path

from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent.parent
EVAL = json.loads((ROOT / '验证结果/材料脱敏评测.json').read_text(encoding='utf-8'))
MODEL = json.loads((ROOT / '验证结果/模型元信息.json').read_text(encoding='utf-8'))
CHECKS = json.loads((ROOT / '验证结果/验证摘要.json').read_text(encoding='utf-8'))
doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Cm(21), Cm(29.7)
section.top_margin, section.bottom_margin = Cm(1.8), Cm(1.8)
section.left_margin, section.right_margin = Cm(2.1), Cm(2.1)
section.footer_distance = Cm(.8)


def font(style, size, east='宋体', bold=False):
    style.font.name = 'Calibri'
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.italic = False
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), east)


font(doc.styles['Normal'], 10.5)
normal = doc.styles['Normal'].paragraph_format
normal.line_spacing = 1.3
normal.space_after = Pt(6)
for name, size in [('Title', 24), ('Subtitle', 13), ('Heading 1', 16), ('Heading 2', 12)]:
    font(doc.styles[name], size, '微软雅黑', name != 'Subtitle')
    doc.styles[name].paragraph_format.space_after = Pt(8)
    doc.styles[name].paragraph_format.space_before = Pt(10 if name.startswith('Heading') else 0)
for name in ['List Bullet', 'List Number']:
    font(doc.styles[name], 10.5)
    doc.styles[name].paragraph_format.line_spacing = 1.25
    doc.styles[name].paragraph_format.space_after = Pt(5)

# Remove the default template's decorative title border.
for style in doc.styles:
    for border in list(style.element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
run = footer.add_run('隐盾整合交付报告  |  ')
run.font.size = Pt(8)
field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE')
footer._p.append(field)
doc.core_properties.title = '隐盾项目完整报告'
doc.core_properties.subject = '材料脱敏与多轮对话整合交付'
doc.core_properties.author = '隐盾项目组'


def p(text, bold_prefix=None):
    paragraph = doc.add_paragraph()
    if bold_prefix and text.startswith(bold_prefix):
        paragraph.add_run(bold_prefix).bold = True
        paragraph.add_run(text[len(bold_prefix):])
    else:
        paragraph.add_run(text)
    return paragraph


def h(text):
    doc.add_heading(text, level=2)


def page(title):
    heading = doc.add_heading(title, level=1)
    heading.paragraph_format.page_break_before = True


def bullet(text):
    doc.add_paragraph(text, style='List Bullet')


def table(headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for column, width in zip(t.columns, widths):
        column.width = Cm(width)
    for cell, text in zip(t.rows[0].cells, headers):
        cell.text = text
    for row in rows:
        for cell, text in zip(t.add_row().cells, row):
            cell.text = str(text)
    for index, row in enumerate(t.rows):
        pr = row._tr.get_or_add_trPr()
        pr.append(OxmlElement('w:cantSplit'))
        if index == 0:
            pr.append(OxmlElement('w:tblHeader'))
        for cell, width in zip(row.cells, widths):
            cell.width = Cm(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tc = cell._tc.get_or_add_tcPr()
            shade = OxmlElement('w:shd')
            shade.set(qn('w:fill'), '29465B' if index == 0 else ('F2F5F7' if index % 2 == 0 else 'FFFFFF'))
            tc.append(shade)
            margins = OxmlElement('w:tcMar')
            for side in ['top','left','bottom','right']:
                edge = OxmlElement(f'w:{side}'); edge.set(qn('w:w'), '90'); edge.set(qn('w:type'), 'dxa'); margins.append(edge)
            tc.append(margins)
            borders = OxmlElement('w:tcBorders')
            for side in ['top','left','bottom','right']:
                edge = OxmlElement(f'w:{side}'); edge.set(qn('w:val'), 'single'); edge.set(qn('w:sz'), '4'); edge.set(qn('w:color'), 'D9D9D9'); borders.append(edge)
            tc.append(borders)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(1)
                paragraph.paragraph_format.line_spacing = 1.15
                for r in paragraph.runs:
                    r.font.size = Pt(9)
                    if index == 0:
                        r.font.bold = True; r.font.color.rgb = RGBColor(255,255,255)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)



from copy import deepcopy
base_doc = deepcopy(doc)
doc.core_properties.subject = '统一对话版本交付与验证'
doc.add_heading('隐盾项目完整报告', 0)
doc.add_paragraph('统一对话版本', style='Subtitle')
p('交付日期  2026 年 10 月 13 日    项目阶段  本机功能验证原型')
doc.add_heading('一 产品目标与本次交付', 1)
p('隐盾提供一个连续对话入口。用户直接输入问题或上传材料，系统在每轮发送前完成本地识别、模型动作决策、内容转换和安全检查，再将处理后的上下文发送给生成服务。材料脱敏是对话内部的处理步骤，文件分析和连续追问共用同一份历史与代号映射。')
p('本次已将首页改为统一对话，移除独立材料处理页面及双入口导航。旧 /chat 地址跳转到首页，不再形成第二套界面。上传材料可在本轮移出、重新加入或从本机删除；处理动作可在对话内展开查看。')
h('用户操作路径')
bullet('打开首页，确认身份和任务，输入消息或添加材料；第一轮发送时自动创建会话。')
bullet('点击发送后，系统自动处理本轮文字与附件，并复检完整上下文。任何片段触发暂停，整轮停止。')
bullet('回答通过检查后显示在同一段对话中。继续补充内容或追问时沿用会话代号，不必重复上传材料。')
bullet('新对话或清除对话会清除当前历史与映射。上传材料独立保存在本机，可通过侧栏删除。')
h('当前交付与验证')
table(['交付项','状态与证据'], [
 ['统一对话界面','已验证上传、两轮对话、附件移出和重加、材料删除、手机宽度布局'],
 ['动作模型与安全检查','由本地模型选动作，安全检查只能放行或暂停'],
 ['API 接入','已有兼容接口客户端；已配置用户指定兼容接口和模型；本机密钥为空，真实服务尚未验收'],
 ['自动化检查','81 项回归、27 个子检查及 38 项冒烟检查通过'],
 ['完整交付','源码、模型、Windows 运行环境、依赖、合成数据、报告与人工复核方案'],
], [4,12.8])
p('当前结论：统一产品流程已在本机跑通。项目仍缺少人工独立金标和真实生成服务验收；自动化通过不等于真实材料准确率、零泄露保证或生产就绪。', '当前结论：')

page('二 本地模型与隐私处理')
h('模型职责')
p('本地动作模型是 TF-IDF 加逻辑回归分类器，负责选择处理动作，不是生成式大语言模型。敏感信息识别使用格式规则、词典和表格上下文；可选校园 NER 默认关闭。远端生成服务负责根据检查后的文字作答，本包不包含云端模型权重。')
p('旧版按披露代价优先排序的规则决策已移除。线上直接采用模型唯一最高概率的动作；安全约束拒绝时暂停，不改选第二名，不回退到规则方案。模型缺失、旧权重、异常概率、并列最高概率或推理失败均停止处理。')
table(['动作','执行含义'], [
 ['DELETE','删除识别出的实体值，尽量保留其他信息'],
 ['LOCAL_AGGREGATE','在本地汇总，仅发送检查通过的统计内容'],
 ['GENERALIZE','降低信息精度，仍受类别、任务和接收方限制'],
 ['PSEUDONYMIZE','对支持的实体替换稳定代号；其他实体按转换器约定删除'],
 ['RETAIN','仅在身份、任务和接收方约束允许时保留原值'],
 ['PAUSE','停止整轮；模型或安全校验均可触发'],
], [4.3,12.5])
h('训练资产与指标含义')
v = MODEL['validation']
p(f'权重位于 app/models/necessity_action.joblib，约 1 MB。从 210 个合成开发案例形成 10,554 条训练样本变体。按来源分组后，训练 {v["train_samples"]:,} 条、{v["train_groups"]} 组；验证 {v["validation_samples"]:,} 条、{v["validation_groups"]} 组；来源组重叠为 0。')
p('当前标签由离线规则启动基线生成，分组留出准确率为 0.9922，Macro-F1 为 0.9887，说明模型能复现这批合成标签；不能将其写成真实场景效果。任务自由描述和历史片段已进入特征，但尚缺针对任务变化、指代和上下文积累的人工验证。')
h('共享处理链路')
p('文字与附件 → 权限检查 → 敏感识别 → 模型选动作 → 动作安全检查 → 本地转换 → 与历史合并 → 完整载荷检查 → 生成服务 → 回答复检 → 提交历史。没有敏感实体也必须确认本地动作模型可用。')
p('已修正整轮暂停、不同接收方白名单混用、识别器异常被当成无敏感信息、实体替换坐标漂移及重叠区间残留等问题。安全检查仍受检测器覆盖范围限制。')

page('三 API 接入状态与运行方式')
h('是否已经接 API')
p('已编写并接入生成服务调用代码，统一对话的发送路径会经过该客户端。用户已提供 https://www.ccsub.net/api/v1 兼容地址和 gpt-5.4 模型名。本机设为 real，但密钥为空；自检未发起网络请求，真实云端服务尚未连接。界面显示状态，不把“已配置”表述成“已连通”。')
table(['配置或结果','实际行为','界面含义'], [
 ['mock','完全离线返回演示内容','未调用真实 API'],
 ['real 且密钥存在','向配置地址发起请求，失败则停止','已配置，发送时连接'],
 ['real 且密钥缺失','停止生成，不降级演示','API 未配置'],
 ['auto','有密钥时尝试真实服务，失败可演示降级','降级须明确标为离线回复'],
 ['收到有效回答且复检通过','写入历史并展示','本轮真实 API 调用成功'],
], [4.2,6.9,5.7])
h('接口调用约定')
p('客户端以 POST 请求发送到 LLM_BASE_URL 加 /chat/completions，使用 Bearer 密钥，载荷含 model、messages、temperature、max_tokens 和 stream=false。messages 包括服务端系统提示、已处理历史与本轮内容；不发送原值映射。仅展示最终 content，不展示 reasoning_content。')
p('已包含 4 项 HTTP 模拟契约检查和 7 项连接自检检查，验证完整 messages、鉴权头、无密钥阻断、错误内容隔离以及纯推理响应拒绝。模拟 HTTP 通过只能证明代码协议行为；真实地址、可用模型、配额、网络与供应商响应仍待实测。')
h('配置和验收步骤')
bullet('双击根目录“启动隐盾.cmd”，访问 http://127.0.0.1:8787/。共享压缩包使用空密钥 mock 配置；本机工作目录保留 real 配置。')
bullet('在本机 app/.env 填写 LLM_API_KEY，确认 LLM_MODE=real 后重启。地址和模型名已填好；双击“检查API连接.cmd”仅发送固定公开短句，真实可用性以返回结果为准。')
bullet('先用不含个人信息的短消息验证连接，再用虚构材料检查两轮上下文。确认界面显示真实 API 成功，同时保留脱敏检查记录。')
bullet('真实服务验收时使用 real，避免自动降级混入结果。缺少真实回复的案例，其最终回答质量应记为未评估。')
p('本地模型仍独立运行，不需要生成服务密钥；生成服务密钥不能提升本地动作模型的训练质量。密钥不应写入报告、截图或共享交付包。')

page('四 会话状态与验证结果')
h('数据保存和生命周期')
p('每轮的文字与附件使用同一会话映射；历史只包含转换后的用户消息及通过复检的回答。用于原值匹配的实体记录和映射仍在本机内存中，因此“历史已脱敏”不等于进程完全不保存原值。失败轮不覆盖此前成功历史。')
p('上传解析结果保存在 app/data/uploads，原始二进制不保留。侧栏删除本机材料只删除该解析记录，不撤回已发送内容或改写历史。清除对话删除历史、映射与原值记录，但不自动删除本机材料。')
table(['约束','当前设置'], [
 ['闲置过期与容量','30 分钟；过期会话在后续访问时清理；最多 64 个会话'],
 ['轮数与文本','最多 20 轮；本轮材料 24000 字符，上下文 64000 字符'],
 ['输入界面限制','消息 12000 字符；每轮最多 10 份附件'],
 ['并发与超限','同一会话一次一轮；超限暂停，不静默裁掉历史'],
 ['恢复能力','刷新页面不自动恢复；服务重启后内存会话丢失'],
], [4.5,12.3])
h('验证记录')
p('本轮使用交付目录中的 Python 运行 81 项回归测试，全部通过，另有 27 个子检查通过。浏览器实测通过统一首页、旧地址跳转、附件管理、两轮脱敏对话、会话清除和手机宽度检查。38 项冒烟检查已重新执行并通过。真实供应商 API 尚未测试。')
h('合成材料基线评测')
rows=[]
for key,label in [('uniform','统一替换'),('fixed','固定类别动作'),('full','当前模型决策')]:
    m=EVAL['summary'][key]
    rows.append([label,f"{m['task_completion_rate']*100:.1f}%",f"{m['block_rate']*100:.1f}%",f"{m['leak_event_rate']*100:.1f}%"])
table(['方案','任务完成代理指标','拦截率','检测口径内泄露事件率'],rows,[3.7,4.2,2.5,6.4])
p('修复姓名误报并仅用开发集重训后，90 例离线基线的模型方案完成代理指标为 90/90，三个任务各 30/30。此前为 82/90；本批已用于开发排查，不是未见验收集。固定类别对照同为 100%，不能据此声称模型更优或真实多轮准确率达 100%。')

page('五 限制与人工复核安排')
h('已知问题')
bullet('“完成海报”中的姓名误报已修复，并增加标签、称谓、复姓、名单和同名重复的回归检查。识别仍依赖启发式，未带标签的人名、复杂指代和生僻姓名需要更多人工样例验证。')
bullet('动作主要按句决策，混合实体容易被统一处理；尚需实体级选择、跨句关系与多轮指代专项评估。')
bullet('数据没有人工独立金标，训练与测试共享合成方式；现有测试已用于排查，不能视为完全未见的最终验收集。')
bullet('人物、事项、期限同时出现不代表对应关系正确；零已检出事件不等于零风险。缺少金标时不报告误拦率。')
bullet('身份下拉框和随机会话 ID 不等于生产认证；多人部署前需补登录、材料归属和会话鉴权。')
h('小规模复核与后续扩充')
p('随包附独立人工复核方案和可填写记录页。先用 3 份材料和 1 组三轮对话试标，统一规则；再锁定 12 份新材料，每份设计两个任务条件，共 24 个请求，另加 6 组三轮对话共 18 轮。正式批次共 42 个请求或轮次。样本是计划数量，尚未完成采集或人工标注。')
p('两名复核者在查看模型结果前分别标出实体、禁止外发值、允许动作集合和必须保留的关系；分歧由已安排的第三人裁决。随后固定模型与配置运行，检查脱敏文本、累计历史和最终回答。')
p('先报告分子和分母，再计算泄露、误放行、误暂停、实体识别和关系保持指标。无真实 API 的回复质量标为未评估；暂停案例单独计入完成率，不得从结果里删掉。完整操作顺序和试用门槛见随包方案。')
h('交付与下一步')
p('根目录包含启动、验证、训练脚本与唯一“打开隐盾.url”；app 包含源码、模型和数据；runtime/python 包含 Windows x64 Python 3.12.6 与依赖；报告和验证结果分别保存文档与证据。已在本机验证，尚未在另一台干净电脑验收。')
p('已准备 9 个合成试标单元和离线复核页，以及两份记录的分歧比较工具。先完成试标，再补独立样本；本机填好密钥后进行真实 API 验收。后续在未见材料与累计对话上验证改进，最后决定是否进入多人试用。')
p('证据位置：验证结果/pytest.xml、网页交互检查.json、材料脱敏评测.json、模型元信息.json 和冒烟测试.txt。此前报告中的 109 项测试、NER F1、FreeAL 与单文件 exe 等历史结论未作为本次验证证据。', '证据位置：')
report_path=ROOT/'报告/隐盾项目完整报告_统一对话版.docx'
doc.save(report_path)
print(report_path)

doc=deepcopy(base_doc)
doc.sections[0].footer.paragraphs[0].runs[0].text='隐盾人工复核方案  |  '
doc.core_properties.title='隐盾人工复核方案'
doc.core_properties.subject='统一对话的标注试行与独立验收'
doc.add_heading('隐盾人工复核方案',0)
doc.add_paragraph('统一对话的标注与验收',style='Subtitle')
p('适用对象  项目开发与复核人员    版本日期  2026 年 10 月 13 日')
doc.add_heading('一 人员与样本安排',1)
p('本方案以少量材料建立可信金标，验证“上传或粘贴 → 自动脱敏 → 连续追问 → 回答”的完整流程。首轮目标是定位错误和判断是否能进入小范围试用，不以小样本推断生产安全性。以下均为拟执行方案，尚未形成已复核数据。')
h('人员分工')
p('复核者 A 和 B 分别独立标注，先不看模型动作与回答。人员安排已确认：A、B 两人独立复核，C 一人裁决分歧并记录依据；不能以模型结果裁决。执行人员固定代码、权重和配置，运行案例、保存输出，不替标注者修改金标。')
h('最小可执行批次')
table(['阶段','材料与任务','对话','用途'],[
 ['试标','3 份材料，每份 2 个任务条件，共 6 个请求','1 组 × 3 轮','修订标注说明，不计最终验收'],
 ['正式复核','12 份新材料，每份 2 个任务条件，共 24 个请求','6 组 × 3 轮，共 18 轮','锁定后验收，共 42 个请求或轮次'],
], [2.2,6.3,3.5,4.8])
p('12 份材料按活动统计、会议纪要、项目介绍各 4 份选取，混合正文、表格和跨句表达。每份材料设置两个有明确差异的任务条件；同一材料及其所有变体属于同一组，不跨训练和验收集。对话材料也应与试标、开发材料独立。')
p('没有足够真实材料时，可以人工新写虚构案例并标为合成来源，再由两人复核；不能把已有训练样本原文改名后称为独立测试。真实材料须经授权并去除不必要的真实身份信息，保留复核所需结构。')
h('六组对话必须覆盖')
bullet('补充同一人物的新任务；新增人物并保持旧代号；代词和省略指代。')
bullet('上传表格后追问汇总；要求恢复姓名或联系方式；多轮分别提供信息导致累计识别风险。')
p('另单列故障检查：模型缺失、API 无密钥或失败、超长上下文、重复提交。故障案例验证停止机制，不与业务准确率混算。含材料准备、标注和裁决的工作量预计 12–18 人时，试标后再调整。分日执行见配套《隐盾项目分工与人工复核执行计划》。')

page('二 金标怎么写')
h('先标预期再看输出')
p('每个案例先记录身份、任务、完整任务要求、接收方和全部原始轮次。当前界面的外发接收方固定为外部生成服务；其他接收方需要单独的接口测试，不能假装已在当前页面覆盖。原始输入保存在本机受控复核目录，不放进公共报告。')
table(['标注对象','必须填写的内容'],[
 ['实体','原文、精确起止位置、类别；重复出现分别标记。按 Unicode 字符从 0 计数，使用左闭右开区间，原文[start:end] 应等于实体值'],
 ['披露限制','本轮及累计历史禁止外发的原值、允许保留的信息和依据'],
 ['允许动作','按片段或实体列出允许动作集合，说明哪些动作虽安全但破坏任务'],
 ['请求结论','是否应放行；如应暂停，写明缺少权限、风险或必要信息'],
 ['任务和关系','必须保留的事实、人物与事项及期限的对应关系、合格回答要点'],
 ['多轮状态','每轮新增信息、指代对象、应沿用的代号、累计禁止信息'],
], [3.2,13.6])
h('动作判定原则')
p('DELETE 用于不应发送且对任务不必要的实体；LOCAL_AGGREGATE 用于只需统计结果的名单；GENERALIZE 用于可降低精度的内容；PSEUDONYMIZE 用于需要保留关系而不需要身份的内容；RETAIN 仅在任务和授权允许时使用；PAUSE 用于无法安全完成或必要条件缺失的请求。不要把固定类别动作当成人工结论。')
p('同一片段可能有多个合格动作，标注允许集合比强行指定一个动作更合理。当前模型按句选择动作；人工可以标注实体级理想结果，评估时分别报告句级实现的限制和可执行金标，不能把粒度差异隐藏掉。')
h('试标示例')
p('虚构输入：“负责人：张明。任务：海报设计，期限：2026年10月20日。电话13812345678。”任务为整理对外模型使用的会议纪要。应遮盖姓名和电话；姓名可用稳定代号，保留事项与期限及其关系。电话不应因姓名采取代号化而残留。')
p('如果“完成海报”中的“成海报”被标为姓名，应记录普通词误报和任务信息损失；如果姓名改了代号但事项对应到另一人，应记录关系错误。即使没有直接原值泄露，也不能判作完全通过。')

page('三 执行顺序与指标')
h('一次完整复核')
bullet('第一步：A、B 独立完成第 4 页记录表的金标部分。比较差异，试标阶段修订说明；正式阶段记录裁决并锁定金标。')
bullet('第二步：记录代码版本、模型 SHA256、配置和运行模式；关闭调参，不使用正式批次重训。每个独立案例新建对话，多轮案例必须连续执行。')
bullet('第三步：逐轮保存原始输入、本轮处理后的文字、完整发送载荷、回答和暂停理由。界面可查看处理后的消息与动作；合成试标工具可导出通过前置校验后交给生成适配器的完整 messages；真实模式才代表尝试网络请求，网页对话入口仍未提供通用载荷导出。')
bullet('第四步：对照金标检查实体边界、动作、全部历史中的禁止值、代号一致性、关系和回答。A、B 独立评分后裁决。供应商请求与日志不保存密钥。')
bullet('第五步：将错误定位到识别、模型决策、转换、网关、历史或生成服务；修改只能在开发数据上进行。正式验收集被用于修复后，要另建未见批次重新验收。')
h('计算口径')
table(['指标','分子与分母'],[
 ['实体识别','严格匹配类别与边界的 TP；Precision=TP/预测实体数，Recall=TP/金标实体数；重复实体按位置计算'],
 ['动作合格率','模型动作属于人工允许集合的片段数 / 已标注且可按当前粒度比较的片段数'],
 ['误放行率','人工判应暂停却实际发送的请求数 / 人工判应暂停的请求数'],
 ['误暂停率','人工判可安全放行却实际暂停的请求数 / 人工判可放行的请求数'],
 ['泄露事件率','实际载荷含禁止值或被人工确认可识别泄露的请求数 / 实际外发请求数；回答泄露另列'],
 ['任务完成与关系','达成金标要点的请求数 / 应可完成的请求数；暂停按未完成。关系正确数 / 金标关系数'],
 ['多轮一致性','代号与指代均正确的跨轮引用数 / 金标跨轮引用数；另列整段对话通过率'],
], [3.7,13.1])
p('分母为 0 时写“不适用”，不要写 100%。未接真实 API 时，真实回答质量写“未评估”；离线回声不计作任务完成。报告原始数量、错误案例和 A/B 裁决前一致率，不能只报平均分。')

page('四 复核记录与试用门槛')
h('单案例记录页')
p('复制本页用于每个独立请求；三轮对话按同一案例 ID 附三份逐轮记录。先填写金标部分，锁定后再填写系统结果。')
p('案例 ID __________    来源组 __________    试标或正式 __________')
p('来源与授权说明 ____________________    复核者 A ____ B ____    日期 __________')
p('身份 __________    任务与完整要求 ____________________    接收方 __________')
p('轮次与原文文件位置 ____________________    前序上下文 ____________________')
h('金标部分')
p('实体原文、类别与位置 _________________________________________________')
p('禁止外发的信息 _______________________________________________________')
p('允许动作集合及理由 ___________________________________________________')
p('应放行或暂停及依据 ___________________________________________________')
p('必须保留的事实与关系 _________________________________________________')
h('系统结果与裁决')
p('代码版本与模型哈希 ____________________    模式 mock 或 real __________')
p('处理文本、发送载荷和回答记录位置 _____________________________________')
p('实际动作与暂停理由 ___________________________________________________')
p('泄露／误暂停／识别／动作／关系／回答错误 ______________________________')
p('A 判断 __________    B 判断 __________    分歧及最终裁决 _______________')
p('修复责任人与回归案例 ID _______________________________________________')
h('首轮内部试用门槛建议')
p('正式 42 个请求或轮次均完成复核，无未解决的高风险分歧；已检出的禁止值泄露和应暂停却发送事件为 0；模型失效仍发送为 0。人工判可放行请求中的误暂停率建议不高于 10%，关键人物与任务关系正确率建议不低于 90%。这些是拟定的内部试用门槛，应在正式执行前确认并锁定，不能事后为通过而修改。')
p('未完成真实 API 验收时，只能通过离线防护流程试用，不能通过真实回答能力验收。任何直接敏感值泄露先修复再测试；小样本无事件只能表述为“本批次未发现”，不能宣称生产零风险。')
plan_path=ROOT/'报告/隐盾人工复核方案与记录表.docx'
doc.save(plan_path)
print(plan_path)
